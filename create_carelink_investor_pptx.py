from pathlib import Path
from math import cos, sin, pi
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree


OUT = "CareLink_Investor_Grade_Presentation.pptx"
ASSETS = Path("assets/images")

W, H = Inches(13.333), Inches(7.5)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
MINT = RGBColor(0x14, 0xB8, 0xA6)
CYAN = RGBColor(0x67, 0xE8, 0xF9)
BLUE = RGBColor(0x60, 0xA5, 0xFA)
EMERALD = RGBColor(0x10, 0xB9, 0x81)
INK = RGBColor(0x0F, 0x17, 0x2A)
SLATE = RGBColor(0x47, 0x55, 0x69)
MUTED = RGBColor(0x64, 0x74, 0x8B)
LINE = RGBColor(0xD7, 0xE5, 0xE8)
PANEL = RGBColor(0xF8, 0xFA, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROSE = RGBColor(0xF4, 0x3F, 0x5E)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
VIOLET = RGBColor(0x8B, 0x5C, 0xF6)

FONT = "Segoe UI"


prs = Presentation()
prs.slide_width = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]


def hexrgb(c):
    return f"{c[0]:02X}{c[1]:02X}{c[2]:02X}"


def solid(shape, color, transparency=0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.fill.transparency = transparency


def no_line(shape):
    shape.line.fill.background()


def line(shape, color=LINE, width=Pt(1)):
    shape.line.color.rgb = color
    shape.line.width = width


def round_rect(slide, x, y, w, h, fill=WHITE, outline=LINE, radius=0.18, transparency=0):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    solid(s, fill, transparency)
    if outline:
        line(s, outline, Pt(1))
    else:
        no_line(s)
    s.adjustments[0] = radius
    return s


def rect(slide, x, y, w, h, fill=WHITE, outline=None, transparency=0):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    solid(s, fill, transparency)
    if outline:
        line(s, outline)
    else:
        no_line(s)
    return s


def oval(slide, x, y, w, h, fill=WHITE, outline=LINE, transparency=0):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    solid(s, fill, transparency)
    if outline:
        line(s, outline, Pt(1))
    else:
        no_line(s)
    return s


def grad(shape, c1, c2, angle=135):
    spPr = shape._element.find(qn("p:spPr"))
    for tag in ("a:solidFill", "a:gradFill", "a:noFill", "a:blipFill", "a:pattFill"):
        for el in spPr.findall(tag.replace("a:", "{http://schemas.openxmlformats.org/drawingml/2006/main}")):
            spPr.remove(el)
    xml = f"""
    <a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
      <a:gsLst>
        <a:gs pos="0"><a:srgbClr val="{hexrgb(c1)}"/></a:gs>
        <a:gs pos="100000"><a:srgbClr val="{hexrgb(c2)}"/></a:gs>
      </a:gsLst>
      <a:lin ang="{int(angle * 60000)}" scaled="0"/>
    </a:gradFill>"""
    spPr.insert(0, etree.fromstring(xml))


def shadow(shape, alpha="18000", blur="65000", dist="22000"):
    spPr = shape._element.find(qn("p:spPr"))
    effect = etree.fromstring(
        f"""
        <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
          <a:outerShdw blurRad="{blur}" dist="{dist}" dir="5400000" algn="ctr" rotWithShape="0">
            <a:srgbClr val="0F172A"><a:alpha val="{alpha}"/></a:srgbClr>
          </a:outerShdw>
        </a:effectLst>"""
    )
    old = spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
    if old is not None:
        spPr.remove(old)
    spPr.append(effect)


def txt(slide, text, x, y, w, h, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def body(slide, text, x, y, w, h, size=12, color=SLATE):
    box = txt(slide, text, x, y, w, h, size, color, False)
    for p in box.text_frame.paragraphs:
        p.line_spacing = 1.05
    return box


def label(slide, s, x=Inches(0.72), y=Inches(0.42)):
    return txt(slide, s.upper(), x, y, Inches(3.0), Inches(0.22), 8, TEAL, True)


def title(slide, s, y=Inches(0.66), x=Inches(0.72), w=Inches(7.8), size=28):
    return txt(slide, s, x, y, w, Inches(0.55), size, INK, True)


def bg(slide, light=True):
    rect(slide, 0, 0, W, H, WHITE if light else TEAL, None)
    a = oval(slide, Inches(9.2), Inches(-1.1), Inches(4.9), Inches(3.1), CYAN, None, 58)
    b = oval(slide, Inches(10.1), Inches(4.9), Inches(4.2), Inches(3.0), MINT, None, 70)
    c = oval(slide, Inches(-1.3), Inches(5.2), Inches(3.2), Inches(2.2), BLUE, None, 82)
    for shp in (a, b, c):
        shp._element.getparent().remove(shp._element)
        slide.shapes._spTree.insert(2, shp._element)
    for col in range(6):
        for row in range(4):
            dot = oval(slide, Inches(12.0 + col * 0.09), Inches(0.35 + row * 0.09), Inches(0.015), Inches(0.015), MINT, None, 15)
            dot._element.getparent().remove(dot._element)
            slide.shapes._spTree.insert(4, dot._element)


def add_transition(slide, kind="fade"):
    tag = "<p:fade/>" if kind == "fade" else "<p:morph/>"
    xml = f'<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med" advClick="1">{tag}</p:transition>'
    slide._element.append(etree.fromstring(xml))


def add_fade_anims(slide, shapes):
    items = ""
    base = 2000 + slide.slide_id * 10
    for i, shp in enumerate(shapes):
        sid = shp.shape_id
        items += f"""
        <p:par><p:cTn id="{base+i}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="clickEffect">
          <p:stCondLst><p:cond delay="{i*80}"/></p:stCondLst>
          <p:childTnLst><p:set><p:cBhvr><p:cTn id="{base+i+1000}" dur="1" fill="hold"/>
          <p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
          </p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set><p:animEffect transition="in" filter="fade"/></p:childTnLst>
        </p:cTn></p:par>"""
    timing = f"""
    <p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
      <p:tnLst><p:par><p:cTn id="{base+5000}" dur="indefinite" restart="whenNotActive" nodeType="tmRoot">
        <p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="{base+5001}" dur="indefinite" nodeType="mainSeq">
          <p:childTnLst>{items}</p:childTnLst></p:cTn>
          <p:prevCondLst><p:cond evt="onPrevClick" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
          <p:nextCondLst><p:cond evt="onNextClick" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
        </p:seq></p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst/>
    </p:timing>"""
    slide._element.append(etree.fromstring(timing))


def icon_circle(slide, x, y, glyph, color=TEAL, size=22):
    c = oval(slide, x, y, Inches(0.55), Inches(0.55), RGBColor(0xEC, 0xFE, 0xFF), RGBColor(0xB6, 0xEE, 0xEA))
    txt(slide, glyph, x + Inches(0.08), y + Inches(0.08), Inches(0.38), Inches(0.30), size, color, False, PP_ALIGN.CENTER)
    return c


def card(slide, x, y, w, h, heading, sub="", glyph=None, accent=TEAL):
    sh = round_rect(slide, x, y, w, h, WHITE, LINE, 0.12)
    shadow(sh, "10000", "42000", "14000")
    if glyph:
        icon_circle(slide, x + Inches(0.22), y + Inches(0.20), glyph, accent, 17)
        tx = x + Inches(0.88)
        tw = w - Inches(1.08)
    else:
        tx = x + Inches(0.24)
        tw = w - Inches(0.48)
    txt(slide, heading, tx, y + Inches(0.20), tw, Inches(0.28), 13, INK, True)
    if sub:
        body(slide, sub, tx, y + Inches(0.55), tw, h - Inches(0.68), 9.5, SLATE)
    return sh


def connector(slide, x1, y1, x2, y2, color=TEAL):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(1.5)
    c.line.end_arrowhead = True
    return c


def phone_mock(slide, x, y, w, h):
    shell = round_rect(slide, x, y, w, h, RGBColor(0x10, 0x13, 0x1A), None, 0.14)
    screen = round_rect(slide, x + Inches(0.10), y + Inches(0.12), w - Inches(0.20), h - Inches(0.24), WHITE, None, 0.12)
    shadow(shell, "26000", "58000", "17000")
    txt(slide, "Find Care", x + Inches(0.27), y + Inches(0.42), w - Inches(0.55), Inches(0.25), 11, INK, True)
    round_rect(slide, x + Inches(0.28), y + Inches(0.83), w - Inches(0.56), Inches(0.35), PANEL, LINE, 0.45)
    txt(slide, "AI Recommendation", x + Inches(0.40), y + Inches(0.91), Inches(1.4), Inches(0.13), 6.8, TEAL, True)
    for i, name in enumerate(("Dr. Ahmed", "Nurse Lina", "Dr. Sara")):
        yy = y + Inches(1.36 + i * 0.76)
        round_rect(slide, x + Inches(0.25), yy, w - Inches(0.50), Inches(0.58), WHITE, LINE, 0.16)
        oval(slide, x + Inches(0.38), yy + Inches(0.12), Inches(0.34), Inches(0.34), RGBColor(0xCC, 0xFB, 0xF1), None)
        txt(slide, name, x + Inches(0.82), yy + Inches(0.12), Inches(1.0), Inches(0.15), 7.2, INK, True)
        txt(slide, "4.9  -  Available", x + Inches(0.82), yy + Inches(0.32), Inches(1.25), Inches(0.13), 6.2, SLATE)
        txt(slide, "0.9", x + w - Inches(0.58), yy + Inches(0.22), Inches(0.24), Inches(0.13), 6.2, TEAL, True)
    return shell


def add_picture(slide, name, x, y, w, h):
    p = ASSETS / name
    if p.exists():
        pic = slide.shapes.add_picture(str(p), x, y, w, h)
        return pic
    return None


slides_anim = []

# 1 Cover
s = prs.slides.add_slide(BLANK); bg(s); add_transition(s, "fade")
logo = add_picture(s, "carelink_logo_dark.png", Inches(0.70), Inches(0.55), Inches(1.85), Inches(0.52))
t1 = txt(s, "AI-Powered\nHome Healthcare,\nRedefined", Inches(0.75), Inches(1.55), Inches(4.8), Inches(1.55), 29, INK, True)
t2 = body(s, "Intelligent care. Trusted providers. Better lives.", Inches(0.78), Inches(3.18), Inches(4.2), Inches(0.35), 12, SLATE)
roles = [("Patient", "PT"), ("Doctor", "DR"), ("Nurse", "NR"), ("Admin", "AD")]
for i, (r, g) in enumerate(roles):
    card(s, Inches(0.75 + i * 1.28), Inches(3.88), Inches(1.12), Inches(0.70), r, "", g, TEAL)
body(s, "Graduation Project\nComputer Science Department\nAcademic Year 2025 - 2026", Inches(0.78), Inches(6.12), Inches(3.9), Inches(0.55), 8.5, INK)
phone_mock(s, Inches(5.95), Inches(0.75), Inches(2.15), Inches(5.95))
pic = add_picture(s, "doctorportrait.jpg", Inches(8.25), Inches(0.82), Inches(3.0), Inches(5.65))
if pic: shadow(pic, "16000", "50000", "14000")
slides_anim.append((s, [t1, t2]))

# 2 Introduction
s = prs.slides.add_slide(BLANK); bg(s); add_transition(s, "fade")
label(s, "Introduction"); title(s, "What is CareLink?")
body(s, "CareLink is a full-stack mobile platform for home healthcare delivery. It connects patients with verified doctors and nurses through intelligent provider matching, real-time scheduling, structured medical records, in-app chat, and governed payments.", Inches(0.75), Inches(1.50), Inches(4.3), Inches(1.25), 12, SLATE)
for i, (h, sub, g) in enumerate([
    ("Patient", "Book care, manage records, chat", "PT"),
    ("Doctor", "Accept requests, schedule, reports", "+"),
    ("Nurse", "Home visits, tracking, earnings", "NR"),
]):
    card(s, Inches(5.55 + i * 2.0), Inches(1.36), Inches(1.62), Inches(2.12), h, sub, g, TEAL)
for i, (n, k) in enumerate([("4", "User roles"), ("12+", "API route groups"), ("AR / EN", "Bilingual interface"), ("AI", "Smart matching")]):
    x = Inches(1.0 + i * 2.95)
    panel = round_rect(s, x, Inches(5.55), Inches(2.45), Inches(0.70), WHITE, LINE, 0.40); shadow(panel, "9000", "36000", "12000")
    txt(s, n, x, Inches(5.72), Inches(2.45), Inches(0.20), 18, TEAL, True, PP_ALIGN.CENTER)
    txt(s, k, x, Inches(6.05), Inches(2.45), Inches(0.16), 7.5, INK, True, PP_ALIGN.CENTER)

# 3 Problem
s = prs.slides.add_slide(BLANK); bg(s); add_transition(s, "fade")
label(s, "The Challenge"); title(s, "Three Gaps in Home Healthcare")
for i, (num, h, sub, col, g) in enumerate([
    ("01", "No Unified Platform", "Patients need one place to discover, verify, book, pay, chat, and manage records.", ROSE, "PT"),
    ("02", "Fragmented Scheduling", "Provider availability, booking requests, rescheduling, and cancellations need coordinated workflows.", AMBER, "CAL"),
    ("03", "Opaque Payments & Governance", "Escrow-style payment status, approval gates, refunds, and wallets require clear control.", VIOLET, "$"),
]):
    x = Inches(0.85 + i * 4.05)
    sh = round_rect(s, x, Inches(1.55), Inches(3.35), Inches(4.35), RGBColor(0xFF, 0xFF, 0xFF), col, 0.08)
    sh.fill.transparency = 6
    shadow(sh, "9000", "42000", "12000")
    txt(s, num, x + Inches(2.55), Inches(1.82), Inches(0.5), Inches(0.28), 16, col, True, PP_ALIGN.RIGHT)
    icon_circle(s, x + Inches(0.35), Inches(2.25), g, col, 21)
    txt(s, h, x + Inches(0.35), Inches(3.25), Inches(2.65), Inches(0.42), 13.5, INK, True)
    body(s, sub, x + Inches(0.35), Inches(3.90), Inches(2.55), Inches(1.15), 10.2, SLATE)

# 4 Solution
s = prs.slides.add_slide(BLANK); bg(s); add_transition(s, "fade")
label(s, "Our Solution"); title(s, "CareLink: End-to-End Home Care")
steps = [("Find\nProviders", "F"), ("Choose\nDate & Time", "CAL"), ("Set\nLocation", "LOC"), ("Secure\nPayment", "$"), ("Care\nDelivered", "OK")]
for i, (h, g) in enumerate(steps):
    x = Inches(0.95 + i * 2.35)
    icon_circle(s, x, Inches(1.80), g, TEAL, 23)
    txt(s, h, x - Inches(0.35), Inches(2.70), Inches(1.25), Inches(0.45), 9.5, INK, True, PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        connector(s, x + Inches(0.63), Inches(2.08), x + Inches(1.52), Inches(2.08))
for i, (h, sub, g) in enumerate([
    ("AI Recommendation", "Smart provider matching", "AI"),
    ("Secure Payments", "Simulated escrow and refund policy", "$"),
    ("Medical Records", "Structured continuous care", "MR"),
    ("Real-time Chat", "Messages, files, voice notes", "CH"),
]):
    card(s, Inches(0.75 + i * 3.12), Inches(4.72), Inches(2.65), Inches(1.05), h, sub, g, TEAL)

# 5 Features
s = prs.slides.add_slide(BLANK); bg(s); add_transition(s, "fade")
label(s, "What We Built"); title(s, "Key Features")
features = [
    ("5-Step Booking Flow", "Search, schedule, location, review, payment", "CAL"),
    ("Real-Time Chat", "Messages, images, PDFs, voice notes", "CH"),
    ("Medical Records", "Diseases, allergies, notes, lab results", "MR"),
    ("Cancellation & Refunds", "Policy-based patient/provider flows", "REF"),
    ("Provider Wallet", "Earnings, payouts, platform commission", "$"),
    ("Rate Approval Gate", "Admin accepts rates before account activation", "OK"),
]
for i, f in enumerate(features):
    card(s, Inches(0.85 + (i % 3) * 4.15), Inches(1.45 + (i // 3) * 2.42), Inches(3.45), Inches(1.72), f[0], f[1], f[2], TEAL)

# 6 AI Engine
s = prs.slides.add_slide(BLANK); bg(s); add_transition(s, "fade")
label(s, "Intelligence Layer"); title(s, "AI Recommendation Engine")
inputs = [("Patient Profile", "diagnoses, allergies, medications", "PT"), ("Medical File Tags", "OCR / AI analysis tags", "MR"), ("Search Query", "free text or voice", "F"), ("GPS Location", "distance and availability", "LOC")]
for i, (h, sub, g) in enumerate(inputs):
    card(s, Inches(0.65 + i * 2.25), Inches(1.45), Inches(1.85), Inches(1.16), h, sub, g, TEAL)
center = oval(s, Inches(5.55), Inches(2.45), Inches(2.2), Inches(2.2), TEAL, None)
grad(center, TEAL, MINT, 135); shadow(center, "22000", "65000", "16000")
txt(s, "AI", Inches(6.10), Inches(2.96), Inches(1.1), Inches(0.52), 31, WHITE, True, PP_ALIGN.CENTER)
txt(s, "matching core", Inches(6.06), Inches(3.48), Inches(1.2), Inches(0.18), 8, WHITE, False, PP_ALIGN.CENTER)
for i in range(4):
    connector(s, Inches(2.45 + i * 2.25), Inches(2.05), Inches(5.55), Inches(3.55), TEAL)
scores = [("Specialization Match", "ICD + keyword"), ("Experience Score", "0.4 - 1.0"), ("Rating Score", "Normalized"), ("Location Score", "<= 1 km = perfect"), ("Medical Compatibility", "Tags vs specialization")]
for i, (h, sub) in enumerate(scores):
    x = Inches(0.72 + i * 2.45)
    card(s, x, Inches(5.24), Inches(2.05), Inches(0.82), h, sub, "*", TEAL)
rank = round_rect(s, Inches(10.20), Inches(1.60), Inches(2.35), Inches(4.35), WHITE, LINE, 0.10); shadow(rank, "14000", "48000", "13000")
txt(s, "Ranked Providers", Inches(10.47), Inches(1.90), Inches(1.8), Inches(0.24), 13, INK, True)
for i, (n, sc) in enumerate([("Dr. Ahmed", "0.92"), ("Nurse Lina", "0.88"), ("Dr. Sara", "0.81")]):
    y = Inches(2.55 + i * 0.82)
    oval(s, Inches(10.50), y, Inches(0.38), Inches(0.38), RGBColor(0xCC, 0xFB, 0xF1), None)
    txt(s, n, Inches(11.02), y + Inches(0.04), Inches(0.9), Inches(0.14), 8.5, INK, True)
    txt(s, sc, Inches(11.95), y + Inches(0.05), Inches(0.32), Inches(0.13), 8, TEAL, True)

# 7 Workflow
s = prs.slides.add_slide(BLANK); bg(s); add_transition(s, "fade")
label(s, "User Journey"); title(s, "Booking Workflow")
flow = [("Provider Discovery", "Patient searches or uses AI"), ("Date & Time", "Availability + blocked slots"), ("Location", "Home address and map context"), ("Review", "Request details and duplicate check"), ("Payment", "Payment status captured"), ("Approval", "Provider accepts and visit starts")]
for i, (h, sub) in enumerate(flow):
    x = Inches(0.75 + i * 2.05)
    card(s, x, Inches(2.05), Inches(1.62), Inches(1.38), h, sub, str(i + 1), TEAL)
    if i < 5:
        connector(s, x + Inches(1.66), Inches(2.72), x + Inches(1.96), Inches(2.72))
round_rect(s, Inches(1.05), Inches(5.05), Inches(11.2), Inches(0.80), RGBColor(0xEC, 0xFE, 0xFF), RGBColor(0xB6, 0xEE, 0xEA), 0.30)
txt(s, "Guardrails: duplicate booking checks - reschedule requests - cancellation/refund policy - notifications", Inches(1.35), Inches(5.32), Inches(10.5), Inches(0.22), 13, TEAL, True, PP_ALIGN.CENTER)

# 8 Architecture
s = prs.slides.add_slide(BLANK); bg(s); add_transition(s, "fade")
label(s, "How It's Built"); title(s, "System Architecture")
cols = [
    ("Clients", ["Patient App", "Doctor App", "Nurse App", "Admin Panel"]),
    ("REST API", ["/patient", "/doctor", "/nurse", "/admin", "/providers", "/payments", "/recommendations", "/chat"]),
    ("Backend Services", ["Node.js", "Express.js", "JWT Auth", "Multer uploads", "Policy services"]),
    ("Database", ["MySQL", "users", "appointments", "payments", "medical records", "wallets"]),
]
for i, (h, rows) in enumerate(cols):
    x = Inches(0.75 + i * 3.15)
    sh = round_rect(s, x, Inches(1.45), Inches(2.42), Inches(4.35), WHITE, LINE, 0.10); shadow(sh, "10000", "42000", "12000")
    txt(s, h, x + Inches(0.22), Inches(1.73), Inches(2.0), Inches(0.20), 11, INK, True, PP_ALIGN.CENTER)
    for j, row in enumerate(rows):
        txt(s, row, x + Inches(0.42), Inches(2.25 + j * 0.43), Inches(1.7), Inches(0.15), 8.5, SLATE)
    if i < 3:
        connector(s, x + Inches(2.46), Inches(3.55), x + Inches(3.05), Inches(3.55))
for i, (h, g) in enumerate([("HTTPS Communication", "API"), ("Auth & Authorization", "OK"), ("File Storage", "FS"), ("Real-time Notifications", "RT")]):
    card(s, Inches(0.90 + i * 3.0), Inches(6.22), Inches(2.36), Inches(0.55), h, "", g, TEAL)

# 9 Stack
s = prs.slides.add_slide(BLANK); bg(s); add_transition(s, "fade")
label(s, "The Stack"); title(s, "Technologies Used")
stack = [("Flutter", "Dart mobile UI"), ("Node.js", "Backend runtime"), ("Express.js", "REST API"), ("MySQL", "Database"), ("JWT", "Authentication"), ("Multer", "File uploads"), ("Firebase FCM", "Planned notifications"), ("GitHub", "Version control")]
for i, (h, sub) in enumerate(stack):
    x = Inches(0.75 + (i % 4) * 3.10)
    y = Inches(1.48 + (i // 4) * 2.30)
    sh = round_rect(s, x, y, Inches(2.52), Inches(1.65), WHITE, LINE, 0.10); shadow(sh, "9000", "40000", "11000")
    txt(s, h[0], x, y + Inches(0.30), Inches(2.52), Inches(0.45), 25, TEAL if i % 2 == 0 else EMERALD, True, PP_ALIGN.CENTER)
    txt(s, h, x, y + Inches(0.93), Inches(2.52), Inches(0.20), 11, INK, True, PP_ALIGN.CENTER)
    txt(s, sub, x, y + Inches(1.22), Inches(2.52), Inches(0.15), 8, SLATE, False, PP_ALIGN.CENTER)

# 10 Roadmap
s = prs.slides.add_slide(BLANK); bg(s); add_transition(s, "fade")
label(s, "What's Next"); title(s, "Future Roadmap")
road = [("2026", "Real Payment Gateway\nStripe / HyperPay", "$"), ("2026+", "Video Consultation\nTelemedicine", "VID"), ("2026+", "Wearables &\nHealth Monitoring", "HW"), ("2026+", "AI Predictive\nHealth Insights", "AI")]
for i, (yr, h, g) in enumerate(road):
    x = Inches(1.00 + i * 3.10)
    icon_circle(s, x, Inches(2.25), g, TEAL, 18)
    if i < 3:
        connector(s, x + Inches(0.56), Inches(2.52), x + Inches(2.50), Inches(2.52))
    txt(s, yr, x - Inches(0.25), Inches(3.35), Inches(1.05), Inches(0.16), 8.5, TEAL, True, PP_ALIGN.CENTER)
    txt(s, h, x - Inches(0.63), Inches(3.65), Inches(1.82), Inches(0.50), 9.4, INK, True, PP_ALIGN.CENTER)

# 11 Conclusion
s = prs.slides.add_slide(BLANK); bg(s); add_transition(s, "fade")
label(s, "Conclusion"); title(s, "CareLink at a Glance")
glance = [("AI-Powered Matching", "Smarter care connections", "AI"), ("Secure Payments", "Governed and protected", "$"), ("Smart Booking", "Real-time conflict-free flow", "CAL"), ("Medical Records", "Structured continuous care", "MR"), ("Better Healthcare", "Accessible, trusted, human", "PT")]
for i, (h, sub, g) in enumerate(glance):
    card(s, Inches(0.75 + i * 2.55), Inches(1.82), Inches(2.05), Inches(1.42), h, sub, g, TEAL)
round_rect(s, Inches(1.25), Inches(5.55), Inches(10.85), Inches(0.55), RGBColor(0xEC, 0xFE, 0xFF), RGBColor(0xB6, 0xEE, 0xEA), 0.35)
txt(s, "CareLink is building the future of home healthcare.", Inches(1.50), Inches(5.73), Inches(10.35), Inches(0.16), 13, TEAL, True, PP_ALIGN.CENTER)

# 12 Thank you
s = prs.slides.add_slide(BLANK); bg(s); add_transition(s, "fade")
add_picture(s, "carelink_logo_dark.png", Inches(0.75), Inches(0.62), Inches(1.95), Inches(0.55))
txt(s, "Thank You!", Inches(0.82), Inches(2.28), Inches(4.4), Inches(0.72), 34, TEAL, True)
txt(s, "AI-Powered Home Healthcare, Redefined", Inches(0.86), Inches(3.22), Inches(4.1), Inches(0.22), 11, TEAL, True)
txt(s, "Any Questions?", Inches(0.86), Inches(4.05), Inches(2.0), Inches(0.18), 10, INK, False)
pic = add_picture(s, "login_brand_illustration.png", Inches(7.20), Inches(0.95), Inches(4.45), Inches(4.90))
if pic: shadow(pic, "12000", "52000", "14000")

for slide in prs.slides:
    try:
        add_fade_anims(slide, list(slide.shapes)[1:8])
    except Exception:
        pass

prs.save(OUT)
print(OUT)
